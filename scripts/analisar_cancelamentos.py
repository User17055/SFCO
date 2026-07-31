import json
import sys
from datetime import date, datetime

from openpyxl import load_workbook


def conv(v):
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    return v


wb = load_workbook(sys.argv[1], data_only=False, read_only=True)
ws = wb["CANCELAMENTOS"]
linhas = []
for numero, row in enumerate(ws.iter_rows(values_only=True), start=1):
    vals = [conv(v) for v in row[:20]]
    if any(v is not None for v in vals):
        linhas.append({"linha": numero, "valores": vals})
    if len(linhas) >= 120:
        break
print(json.dumps({"dimensao": [ws.max_row, ws.max_column], "linhas": linhas}, ensure_ascii=False, indent=2))
