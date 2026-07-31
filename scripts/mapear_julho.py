import collections
import json
import re
import sys
import unicodedata
from datetime import date, datetime

from openpyxl import load_workbook


def limpar(v):
    return " ".join(str(v or "").strip().split())


def sem_acento(v):
    return "".join(
        c for c in unicodedata.normalize("NFD", limpar(v).upper())
        if unicodedata.category(c) != "Mn"
    )


def normalizar_plano(v):
    p = sem_acento(v)
    p = re.sub(r"\s+", " ", p).strip()
    aliases = {
        "SUPER ECONOMICO": "PLANO SUPER ECONOMICO",
        "ECONOMICO FELINO": "PLANO ECONOMICO FELINO",
        "ECONOMICO": "PLANO ECONOMICO",
        "HIGIENICO": "PLANO HIGIENICO",
        "PUPPY": "PLANO PUPPY",
        "PUPY": "PLANO PUPPY",
    }
    return aliases.get(p, p)


wb = load_workbook(sys.argv[1], data_only=True, read_only=True)
ws = wb["JULHO"]
raw = collections.Counter()
normalizados = collections.Counter()
valores = collections.defaultdict(collections.Counter)
invalidos = []
linhas = []

for numero, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    inicio, reajuste, cliente, plano, adicional, valor, pet, pago, meses, comprovante = list(row[:10])
    cliente = limpar(cliente)
    pet = limpar(pet)
    plano_raw = limpar(plano)
    plano_norm = normalizar_plano(plano)
    if not cliente or not pet or not plano_norm or not isinstance(valor, (int, float)):
        invalidos.append({"linha": numero, "cliente": cliente, "pet": pet, "plano": plano_raw, "valor": valor})
        continue
    raw[sem_acento(plano_raw)] += 1
    normalizados[plano_norm] += 1
    valores[plano_norm][round(float(valor), 2)] += 1
    linhas.append({
        "linha": numero,
        "cliente": cliente,
        "pet": pet,
        "plano": plano_norm,
        "valor": round(float(valor), 2),
        "adicional": limpar(adicional),
        "data_inicio": inicio.isoformat()[:10] if isinstance(inicio, (date, datetime)) else limpar(inicio),
        "data_reajuste": reajuste.isoformat()[:10] if isinstance(reajuste, (date, datetime)) else limpar(reajuste),
        "observacao": " | ".join(v for v in [limpar(pago), limpar(meses), limpar(comprovante)] if v),
    })

resultado = {
    "registros_validos": len(linhas),
    "registros_invalidos": len(invalidos),
    "clientes": len({sem_acento(x["cliente"]) for x in linhas}),
    "pets": len({(sem_acento(x["cliente"]), sem_acento(x["pet"])) for x in linhas}),
    "planos_raw": raw.most_common(),
    "planos_normalizados": normalizados.most_common(),
    "valores_por_plano": {k: v.most_common() for k, v in valores.items()},
    "invalidos": invalidos,
    "projecao_total": round(sum(x["valor"] for x in linhas), 2),
}
print(json.dumps(resultado, ensure_ascii=False, indent=2))
