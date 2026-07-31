import collections
import json
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def texto(valor):
    if valor is None:
        return ""
    return " ".join(str(valor).strip().split())


def chave(valor):
    valor = texto(valor).upper()
    valor = "".join(
        c for c in unicodedata.normalize("NFD", valor)
        if unicodedata.category(c) != "Mn"
    )
    return valor


arquivo = Path(sys.argv[1])
wb = load_workbook(arquivo, data_only=True, read_only=True)
resultado = {"arquivo": str(arquivo), "abas": []}

for ws in wb.worksheets:
    cabecalho = [texto(v) for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    planos = collections.Counter()
    clientes = set()
    pets = set()
    linhas_validas = 0
    datas_min = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        # Estrutura principal das abas mensais: A data, C cliente, D plano, E valor, F/G pet.
        cliente = texto(row[2] if len(row) > 2 else None)
        plano = texto(row[3] if len(row) > 3 else None)
        valor = row[4] if len(row) > 4 else None
        pet = texto(row[5] if len(row) > 5 else None)
        if cliente and plano and (pet or isinstance(valor, (int, float))):
            linhas_validas += 1
            clientes.add(chave(cliente))
            if pet:
                pets.add((chave(cliente), chave(pet)))
            planos[chave(plano)] += 1
        data_inicio = row[0] if row else None
        if isinstance(data_inicio, (date, datetime)):
            datas_min.append(data_inicio.isoformat())

    resultado["abas"].append({
        "nome": ws.title,
        "estado": ws.sheet_state,
        "linhas": ws.max_row,
        "colunas": ws.max_column,
        "cabecalho": cabecalho[:20],
        "linhas_validas_padrao": linhas_validas,
        "clientes_distintos": len(clientes),
        "pets_distintos": len(pets),
        "planos_distintos": len(planos),
        "planos_mais_comuns": planos.most_common(12),
        "menor_data": min(datas_min) if datas_min else None,
        "maior_data": max(datas_min) if datas_min else None,
    })

print(json.dumps(resultado, ensure_ascii=False, indent=2))
