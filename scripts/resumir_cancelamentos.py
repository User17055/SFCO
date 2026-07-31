import collections
import json
import sys

from openpyxl import load_workbook


def txt(v):
    return " ".join(str(v or "").strip().upper().split())


wb = load_workbook(sys.argv[1], data_only=True, read_only=True)
ws = wb["CANCELAMENTOS"]
meses = collections.Counter()
motivos = collections.Counter()
registros = []
for numero, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    nome = txt(row[0] if len(row) > 0 else None)
    plano = txt(row[1] if len(row) > 1 else None)
    quando = txt(row[3] if len(row) > 3 else None)
    motivo = txt(row[4] if len(row) > 4 else None)
    valor = row[5] if len(row) > 5 else None
    quantidade = row[6] if len(row) > 6 else None
    if not nome or not quando or not isinstance(quantidade, (int, float)):
        continue
    meses[quando] += int(quantidade)
    motivos[motivo] += int(quantidade)
    registros.append({"linha": numero, "nome": nome, "plano": plano, "mes": quando, "motivo": motivo, "valor": valor, "quantidade": quantidade})

print(json.dumps({
    "registros": len(registros),
    "quantidade_total": sum(int(x["quantidade"]) for x in registros),
    "meses": meses,
    "motivos_principais": motivos.most_common(30),
    "ultimos_30": registros[-30:],
}, ensure_ascii=False, indent=2, default=dict))
