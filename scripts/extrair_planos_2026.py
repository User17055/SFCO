"""Extrai a planilha PLANOS 2026 para um JSON idempotente de importacao."""

from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


NUMERO_MES = {
    "JANEIRO": 1,
    "FEVEREIRO": 2,
    "MARCO": 3,
    "ABRIL": 4,
    "MAIO": 5,
    "JUNHO": 6,
    "JULHO": 7,
    "AGOSTO": 8,
    "SETEMBRO": 9,
    "OUTUBRO": 10,
    "NOVEMBRO": 11,
    "DEZEMBRO": 12,
}


def texto(valor) -> str:
    return " ".join(str(valor or "").strip().split())


def chave(valor) -> str:
    sem_acento = "".join(
        caractere
        for caractere in unicodedata.normalize("NFD", texto(valor).upper())
        if unicodedata.category(caractere) != "Mn"
    )
    return re.sub(r"\s+", " ", sem_acento).strip()


def hash_codigo(*partes) -> str:
    conteudo = "|".join(chave(parte) for parte in partes)
    return hashlib.sha256(conteudo.encode("utf-8")).hexdigest()


def data_iso(valor):
    if isinstance(valor, (date, datetime)):
        return valor.isoformat()[:10]
    valor_texto = texto(valor)
    return valor_texto if re.fullmatch(r"\d{4}-\d{2}-\d{2}", valor_texto) else None


def numero(valor, padrao=0.0):
    if isinstance(valor, (int, float)):
        return round(float(valor), 2)
    return padrao


def inteiro(valor):
    if isinstance(valor, (int, float)):
        return int(valor)
    return None


def normalizar_plano(valor) -> str:
    plano = chave(valor)
    aliases = {
        "SUPER ECONOMICO": "PLANO SUPER ECONOMICO",
        "ECONOMICO FELINO": "PLANO ECONOMICO FELINO",
        "ECONOMICO": "PLANO ECONOMICO",
        "HIGIENICO": "PLANO HIGIENICO",
        "PUPPY": "PLANO PUPPY",
        "PUPY": "PLANO PUPPY",
        "PACOET BANHO QUINZENAL": "PACOTE BANHO QUINZENAL",
        "VIP": "PLANO VIP",
        "3 SUPER ECONOMICO FELINO": "PLANO ECONOMICO FELINO",
        "SUPER ECONOMICO+2 BANHOS": "SUPER ECONOMICO + BANHOS",
    }
    return aliases.get(plano, plano) or "PLANO NAO INFORMADO"


def observacao(*valores) -> str:
    return " | ".join(item for item in (texto(v) for v in valores) if item)


def extrair(caminho: Path) -> dict:
    workbook = load_workbook(caminho, data_only=True, read_only=True)
    ano_encontrado = re.search(r"\b(20\d{2})\b", caminho.stem)
    ano_atual = int(ano_encontrado.group(1)) if ano_encontrado else date.today().year
    abas_mensais = {
        NUMERO_MES[chave(nome)]: nome
        for nome in workbook.sheetnames
        if chave(nome) in NUMERO_MES
    }
    if not abas_mensais:
        raise ValueError("Nenhuma aba mensal foi encontrada na planilha.")
    mes_atual = max(abas_mensais)
    nome_aba_atual = abas_mensais[mes_atual]
    competencia_atual = f"{ano_atual:04d}-{mes_atual:02d}-01"
    historico = []
    atuais = []
    cancelamentos = []
    avisos = []
    valores_planos = defaultdict(Counter)

    detalhes = workbook["DETAILS"]
    for linha, valores in enumerate(detalhes.iter_rows(min_row=2, values_only=True), start=2):
        dados = list(valores[:10]) + [None] * max(0, 10 - len(valores))
        inicio, reajuste, cliente, plano, adicional, qtd_banho, valor_banho, valor, pet, mes = dados[:10]
        cliente_nome = texto(cliente)
        pet_nome = texto(pet)
        mes_chave = chave(mes)
        ano_no_mes = re.search(r"\b(20\d{2})\b", mes_chave)
        mes_sem_ano = re.sub(r"\s+20\d{2}\b", "", mes_chave).strip()
        mes_numero = NUMERO_MES.get(mes_sem_ano)
        ano_competencia = (
            int(ano_no_mes.group(1))
            if ano_no_mes
            else (ano_atual if mes_numero and mes_numero <= mes_atual else ano_atual - 1)
        )
        competencia = (
            f"{ano_competencia:04d}-{mes_numero:02d}-01"
            if mes_numero
            else None
        )
        if not competencia or not cliente_nome:
            continue
        if not pet_nome:
            pet_nome = "PET NAO INFORMADO"
            avisos.append({
                "aba": "DETAILS",
                "linha": linha,
                "cliente": cliente_nome,
                "pet": pet_nome,
                "problema": "pet ausente",
            })
        plano_nome = normalizar_plano(plano)
        valor_mensal = numero(valor)
        historico.append({
            "codigo_externo": hash_codigo("DETAILS", linha),
            "competencia": competencia,
            "cliente_nome": cliente_nome,
            "pet_nome": pet_nome,
            "plano_nome": plano_nome,
            "valor_mensal": valor_mensal,
            "data_inicio": data_iso(inicio),
            "data_reajuste": data_iso(reajuste),
            "adicional": texto(adicional) or None,
            "quantidade_banho": inteiro(qtd_banho),
            "valor_banho": numero(valor_banho, None),
            "observacoes": None,
            "origem": "PLANOS 2026 - HISTORICO",
            "linha_origem": linha,
        })

    # DETAILS nem sempre acompanha as ultimas abas mensais. Completa o historico
    # com cada aba anterior a atual que ainda nao exista em DETAILS. Assim o
    # painel nao compara setembro diretamente com junho, por exemplo.
    competencias_historico = {item["competencia"] for item in historico}
    for mes_numero, nome_aba in sorted(abas_mensais.items()):
        if mes_numero >= mes_atual:
            continue
        competencia = f"{ano_atual:04d}-{mes_numero:02d}-01"
        if competencia in competencias_historico:
            continue
        aba_mensal = workbook[nome_aba]
        for linha, valores in enumerate(aba_mensal.iter_rows(min_row=2, values_only=True), start=2):
            dados = list(valores[:10]) + [None] * max(0, 10 - len(valores))
            inicio, reajuste, cliente, plano, adicional, valor, pet, pago, meses, comprovante = dados[:10]
            cliente_nome = texto(cliente)
            pet_nome = texto(pet)
            if not cliente_nome or not pet_nome:
                continue
            historico.append({
                "codigo_externo": hash_codigo("MENSAL", nome_aba, linha),
                "competencia": competencia,
                "cliente_nome": cliente_nome,
                "pet_nome": pet_nome,
                "plano_nome": normalizar_plano(plano),
                "valor_mensal": numero(valor),
                "data_inicio": data_iso(inicio),
                "data_reajuste": data_iso(reajuste),
                "adicional": texto(adicional) or None,
                "quantidade_banho": None,
                "valor_banho": None,
                "observacoes": observacao(pago, meses, comprovante) or None,
                "origem": "PLANOS 2026 - HISTORICO",
                "linha_origem": linha,
            })
        competencias_historico.add(competencia)

    aba_atual = workbook[nome_aba_atual]
    for linha, valores in enumerate(aba_atual.iter_rows(min_row=2, values_only=True), start=2):
        dados = list(valores[:10]) + [None] * max(0, 10 - len(valores))
        inicio, reajuste, cliente, plano, adicional, valor, pet, pago, meses, comprovante = dados[:10]
        cliente_nome = texto(cliente)
        pet_nome = texto(pet)
        if not cliente_nome or not pet_nome:
            continue
        plano_nome = normalizar_plano(plano)
        valor_mensal = numero(valor)
        if not texto(plano) or not isinstance(valor, (int, float)):
            avisos.append({
                "aba": nome_aba_atual,
                "linha": linha,
                "cliente": cliente_nome,
                "pet": pet_nome,
                "problema": "plano ausente" if not texto(plano) else "valor ausente",
            })
            continue
        atuais.append({
            "codigo_externo": hash_codigo("ATUAL", linha),
            "cliente_codigo": hash_codigo("CLIENTE", cliente_nome),
            "pet_codigo": hash_codigo("PET", cliente_nome, pet_nome),
            "cliente_nome": cliente_nome,
            "pet_nome": pet_nome,
            "plano_nome": plano_nome,
            "valor_mensal": valor_mensal,
            "data_inicio": data_iso(inicio),
            "data_reajuste": data_iso(reajuste),
            "adicional": texto(adicional) or None,
            "observacoes": observacao(pago, meses, comprovante) or None,
            "origem": "PLANOS 2026 - ATUAL",
        })
        if valor_mensal > 0:
            valores_planos[plano_nome][valor_mensal] += 1

    aba_cancelamentos = workbook["CANCELAMENTOS"]
    for linha, valores in enumerate(aba_cancelamentos.iter_rows(min_row=2, values_only=True), start=2):
        dados = list(valores[:7]) + [None] * max(0, 7 - len(valores))
        cliente, plano, tentativa, mes, motivo, valor, quantidade = dados[:7]
        cliente_nome = texto(cliente)
        mes_nome = chave(mes)
        qtd = inteiro(quantidade)
        if not cliente_nome or mes_nome not in NUMERO_MES or qtd is None:
            continue
        mes_numero = NUMERO_MES[mes_nome]
        # Meses ate a aba atual pertencem ao ano da planilha; meses futuros
        # pertencem ao historico do ano anterior.
        ano = ano_atual if mes_numero <= mes_atual else ano_atual - 1
        cancelamentos.append({
            "codigo_externo": hash_codigo("CANCELAMENTOS", linha),
            "cliente_nome": cliente_nome,
            "plano_nome": normalizar_plano(plano) if texto(plano) else None,
            "competencia": f"{ano:04d}-{mes_numero:02d}-01",
            "motivo": texto(motivo) or None,
            "tentativa_recuperacao": texto(tentativa) or None,
            "valor": numero(valor),
            "quantidade": max(qtd, 1),
            "origem": "PLANOS 2026 - CANCELAMENTOS",
            "linha_origem": linha,
        })

    catalogo = []
    for nome, contagem in sorted(valores_planos.items()):
        valor_padrao = contagem.most_common(1)[0][0] if contagem else 0
        catalogo.append({"nome": nome, "valor_padrao": valor_padrao})

    return {
        "arquivo_origem": str(caminho),
        "competencia_atual": competencia_atual,
        "aba_atual": nome_aba_atual,
        "origem_atual": "PLANOS 2026 - ATUAL",
        "historico": historico,
        "atuais": atuais,
        "cancelamentos": cancelamentos,
        "catalogo": catalogo,
        "avisos": avisos,
        "resumo": {
            "historico": len(historico),
            "atuais": len(atuais),
            "clientes_atuais": len({x["cliente_codigo"] for x in atuais}),
            "pets_atuais": len({x["pet_codigo"] for x in atuais}),
            "cancelamentos": len(cancelamentos),
            "tipos_planos": len(catalogo),
            "projecao_atual": round(sum(x["valor_mensal"] for x in atuais), 2),
        },
    }


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Uso: python extrair_planos_2026.py PLANILHA.xlsx SAIDA.json")
    origem = Path(sys.argv[1])
    destino = Path(sys.argv[2])
    resultado = extrair(origem)
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(json.dumps(resultado, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(resultado["resumo"], ensure_ascii=False, indent=2))
    if resultado["avisos"]:
        print("Avisos:")
        print(json.dumps(resultado["avisos"], ensure_ascii=False, indent=2))
