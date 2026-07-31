import collections
import json
import sys
from datetime import date, datetime

from openpyxl import load_workbook


def val(v):
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    return v


wb = load_workbook(sys.argv[1], data_only=True, read_only=True)
ws = wb["DETAILS"]
meses = collections.Counter()
planos_julho = collections.Counter()
valores_planos = collections.defaultdict(collections.Counter)
amostra_julho = []

for row in ws.iter_rows(min_row=2, values_only=True):
    valores = list(row[:10])
    mes = str(valores[9] or "").strip().upper()
    if not mes:
        continue
    meses[mes] += 1
    if mes == "JULHO":
        plano = " ".join(str(valores[3] or "").strip().upper().split())
        planos_julho[plano] += 1
        valor = valores[7]
        if isinstance(valor, (int, float)):
            valores_planos[plano][round(float(valor), 2)] += 1
        if len(amostra_julho) < 30:
            amostra_julho.append([val(v) for v in valores])

resultado = {
    "cabecalho": [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))[:10]],
    "contagem_meses": meses,
    "linhas_julho": meses["JULHO"],
    "planos_julho": planos_julho,
    "valores_mais_comuns_por_plano": {
        plano: contagem.most_common(8)
        for plano, contagem in valores_planos.items()
    },
    "amostra_julho": amostra_julho,
}
print(json.dumps(resultado, ensure_ascii=False, indent=2, default=dict))
