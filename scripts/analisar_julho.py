import collections
import json
import sys
from datetime import date, datetime

from openpyxl import load_workbook


def val(v):
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    return v


wb = load_workbook(sys.argv[1], data_only=True, read_only=True)
ws = wb["JULHO"]
amostra = []
tipos = [collections.Counter() for _ in range(10)]
for row in ws.iter_rows(min_row=2, values_only=True):
    valores = list(row[:10])
    if not any(v is not None for v in valores):
        continue
    for i, v in enumerate(valores):
        tipos[i][type(v).__name__] += 1
    if len(amostra) < 40:
        amostra.append([val(v) for v in valores])

print(json.dumps({
    "cabecalho": [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))[:10]],
    "tipos_por_coluna": [dict(c) for c in tipos],
    "amostra": amostra,
}, ensure_ascii=False, indent=2))
