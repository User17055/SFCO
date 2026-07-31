import json
import sys
from pathlib import Path

from openpyxl import load_workbook


arquivo = Path(sys.argv[1])
workbook = load_workbook(arquivo, data_only=False, read_only=False)
resumo = {"arquivo": str(arquivo), "abas": []}

for aba in workbook.worksheets:
    linhas = []
    formulas = []
    for row in aba.iter_rows(
        min_row=1,
        max_row=min(aba.max_row, 80),
        min_col=1,
        max_col=min(aba.max_column, 30),
    ):
        valores = []
        tem_valor = False
        for cell in row:
            valor = cell.value
            if valor is not None:
                tem_valor = True
            if isinstance(valor, str) and valor.startswith("="):
                formulas.append({"celula": cell.coordinate, "formula": valor})
            if hasattr(valor, "isoformat"):
                valor = valor.isoformat()
            valores.append(valor)
        if tem_valor:
            linhas.append({"linha": row[0].row, "valores": valores})
        if len(linhas) >= 35:
            break

    resumo["abas"].append(
        {
            "nome": aba.title,
            "estado": aba.sheet_state,
            "dimensao": aba.calculate_dimension(),
            "max_linhas": aba.max_row,
            "max_colunas": aba.max_column,
            "mescladas": [str(item) for item in list(aba.merged_cells.ranges)[:30]],
            "tabelas": list(aba.tables.keys()),
            "imagens": len(aba._images),
            "graficos": len(aba._charts),
            "amostra": linhas,
            "formulas_amostra": formulas[:30],
        }
    )

print(json.dumps(resumo, ensure_ascii=False, indent=2))
